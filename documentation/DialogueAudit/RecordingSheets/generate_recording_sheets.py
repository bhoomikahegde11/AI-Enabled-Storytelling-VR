"""Regenerate the six source-faithful Level 1 recording plans and XLSX workbooks.

Requires Python with openpyxl for reading existing recording statuses, Node.js, and
@oai/artifact-tool. Set RECORDING_NODE_MODULES if the bundled runtime is elsewhere.
No Unity files or dialogue inventory files are written.
"""

import csv
import json
import os
import re
import subprocess
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

import sys

AUDIT_DIR = Path(__file__).resolve().parent.parent
sys.dont_write_bytecode = True
sys.path.insert(0, str(AUDIT_DIR))
from extract_level1_dialogues import CHARACTER_FILES, parse_character_file  # noqa: E402

OUT = Path(__file__).resolve().parent
CSV = AUDIT_DIR / "Level1_AllNPC_Dialogues.csv"
PLACEHOLDER = re.compile(r"(\{[A-Za-z][A-Za-z0-9]*\})")
EXPECTED_COUNTS = {"LAK": 206, "CHI": 167, "SAR": 164, "ABD": 170, "FRA": 170, "FAT": 170}
WORKBOOK_NAMES = {
    "LAK": "LakshmiAmma_RecordingSheet.xlsx",
    "CHI": "ChinnammaNaik_RecordingSheet.xlsx",
    "SAR": "SaraswatiChetti_RecordingSheet.xlsx",
    "ABD": "AbdulRahman_RecordingSheet.xlsx",
    "FRA": "Francisco_RecordingSheet.xlsx",
    "FAT": "FatherPenteado_RecordingSheet.xlsx",
}
VOICE = {
    "LAK": "Warm, familiar, quietly shrewd",
    "CHI": "Direct, confident, dry-witted",
    "SAR": "Measured, observant, precise",
    "ABD": "Courteous, sociable, persuasive",
    "FRA": "Expressive, energetic, commercially alert",
    "FAT": "Thoughtful, restrained, gently curious",
}
NUMBER_WORDS = {
    1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six",
    7: "seven", 8: "eight", 9: "nine", 10: "ten", 11: "eleven",
    12: "twelve", 13: "thirteen", 14: "fourteen", 15: "fifteen",
    16: "sixteen", 17: "seventeen", 18: "eighteen", 19: "nineteen",
    20: "twenty", 30: "thirty", 40: "forty", 50: "fifty",
    60: "sixty", 70: "seventy", 80: "eighty", 90: "ninety",
    **{100 * i: f"{['', 'one', 'two', 'three', 'four', 'five', 'six', 'seven', 'eight', 'nine'][i]} hundred" for i in range(1, 10)},
}
NUMBER_FIELDS = {"counterPrice", "currentBuyerOffer", "offeredPrice", "finalPrice"}
SPICES = ["pepper", "cinnamon", "cardamom", "clove"]
UNITS = ["seer", "seers", "veesai", "palam", "palams", "manangu", "manangus", "bahar", "bahars"]

OVERVIEW = ["Dialogue_ID", "NPC", "Scenario", "Mood", "Patience", "Original_Text",
            "Dynamic_Variables", "Segment_Count", "Ordered_Assembly_Sequence", "Status", "Notes"]
STATIC = ["Segment_ID", "Dialogue_ID", "Segment_Order", "Text_To_Record",
          "Suggested_Filename", "Voice_Direction", "Scenario", "Recording_Status", "Notes"]
NUMBERS = ["Number", "Spoken_Text", "Suggested_Filename", "Recording_Status", "Notes"]
VOCAB = ["Vocabulary_ID", "Category", "Spoken_Text", "Suggested_Filename",
         "Used_By_Placeholder", "Recording_Status", "Notes"]
MANIFEST = ["Dialogue_ID", "Segment_Order", "Component_Type", "Source_Text_Or_Placeholder",
            "Playback_Reference", "Runtime_Resolution_Rule", "Is_Required", "Notes"]
SUMMARY = ["Metric", "Value", "Notes"]


def old_rows(book_path, tab, key):
    if not book_path.exists():
        return {}
    wb = load_workbook(book_path, read_only=True, data_only=True)
    try:
        if tab not in wb.sheetnames:
            return {}
        rows = wb[tab].iter_rows(values_only=True)
        headers = next(rows)
        if key not in headers:
            return {}
        return {str(row[headers.index(key)]): dict(zip(headers, row)) for row in rows
                if row[headers.index(key)] is not None}
    finally:
        wb.close()


def old_manifest(book_path):
    if not book_path.exists():
        return {}
    wb = load_workbook(book_path, read_only=True, data_only=True)
    try:
        if "Assembly Manifest" not in wb.sheetnames:
            return {}
        cells = wb["Assembly Manifest"].iter_rows(values_only=True)
        headers = next(cells)
        grouped = defaultdict(list)
        for values in cells:
            row = dict(zip(headers, values))
            grouped[row["Dialogue_ID"]].append(row)
        return grouped
    finally:
        wb.close()


def keep_status(old, expected, status_field):
    if old and all(old.get(k) == v for k, v in expected.items()):
        return old.get(status_field) or "PENDING"
    return "PENDING"


def inventory_from_source():
    source = []
    for name, char_id, filename, prefix in CHARACTER_FILES:
        seq = 0
        for entry in parse_character_file(name, char_id, filename, prefix):
            for text, line in entry["templates"]:
                seq += 1
                source.append({
                    "Dialogue_ID": f"L1_{prefix}_{seq:03d}", "NPC": name,
                    "Scenario": entry["scenario"], "Mood": entry["mood_bucket"],
                    "Patience": entry["patience_bucket"], "Original_Text": text,
                    "Source_File": entry["rel_path"], "Source_Line": line,
                    "Prefix": prefix.lower(), "Prefix_Upper": prefix,
                })
        if seq != EXPECTED_COUNTS[prefix]:
            raise ValueError(f"Unexpected {name} count {seq}; inspect source before regenerating")
    with CSV.open(encoding="utf-8-sig", newline="") as file:
        audit = list(csv.DictReader(file))
    if len(audit) != len(source):
        raise ValueError(f"Audit CSV has {len(audit)} rows; C# has {len(source)}")
    for i, (a, s) in enumerate(zip(audit, source), 1):
        if (a["Dialogue_ID"], a["NPC"], a["Scenario"], a["Mood"], a["Patience"], a["Dialogue_Text"]) != (
                s["Dialogue_ID"], s["NPC"], s["Scenario"], s["Mood"], s["Patience"], s["Original_Text"]):
            raise ValueError(f"CSV/source mismatch at row {i}: {s['Dialogue_ID']}; audit before generating")
    ids = [row["Dialogue_ID"] for row in source]
    if len(ids) != len(set(ids)) or len(source) != 1047:
        raise ValueError("Duplicate IDs or unexpected Level 1 inventory count")
    return source


def vocabulary(prefix):
    rows = []
    def add(category, key, spoken, placeholder, note=""):
        rows.append({"Vocabulary_ID": f"{prefix.upper()}_{category.upper()}_{key.upper()}",
                     "Category": category.upper(), "Spoken_Text": spoken,
                     "Suggested_Filename": f"{prefix}_{category}_{key}.wav",
                     "Used_By_Placeholder": placeholder, "Recording_Status": "PENDING", "Notes": note})
    for spice in SPICES:
        note = "Market catalog only; current local session generator excludes clove." if spice == "clove" else "Current local session spice."
        add("spice", spice, spice, "{spiceName}", note)
    add("spice", "generic", "spice", "{spiceName}", "Provider fallback only.")
    add("currency", "varahas", "varahas", "{currency}", "The current provider always emits this plural form.")
    add("currency", "varaha", "varaha", "{currency}", "Singular form reserved; current provider does not emit it.")
    for unit in UNITS:
        add("unit", unit, unit.capitalize(), "{quantityLabel}", "From MarketManager.FormatTraditionalQuantity.")
    add("quantity", "this_lot", "this lot", "{quantityLabel}", "Provider fallback for absent quantity label.")
    return rows


def direction(prefix, scenario, mood, patience):
    cues = [VOICE[prefix.upper()]]
    if mood != "Any":
        cues.append(f"{mood.lower()} delivery")
    if patience == "Low":
        cues.append("brief; low patience")
    if scenario in {"TransactionFailure", "HostileWarning"}:
        cues.append("controlled firmness")
    elif scenario == "TransactionSuccess":
        cues.append("completed-trade satisfaction")
    elif scenario in {"SellerPriceAccepted", "PlayerAcceptedDeal"}:
        cues.append("price agreement, not delivery")
    return "; ".join(cues)


def build_character(records):
    prefix = records[0]["Prefix"]
    book = OUT / WORKBOOK_NAMES[records[0]["Prefix_Upper"]]
    prior_overview = old_rows(book, "Dialogue Overview", "Dialogue_ID")
    prior_static = old_rows(book, "Static Recording Segments", "Segment_ID")
    prior_numbers = old_rows(book, "Number Library", "Number")
    prior_vocab = old_rows(book, "Dynamic Vocabulary", "Vocabulary_ID")
    prior_manifest = old_manifest(book)
    overview, static, manifest, unresolved = [], [], [], []
    vocab = vocabulary(prefix)
    for row in vocab:
        old = prior_vocab.get(row["Vocabulary_ID"])
        row["Recording_Status"] = keep_status(old, {k: row[k] for k in
            ("Category", "Spoken_Text", "Suggested_Filename", "Used_By_Placeholder")}, "Recording_Status")
    vocab_files = {r["Suggested_Filename"] for r in vocab}
    units = [f"{prefix}_unit_{unit}.wav" for unit in UNITS]
    spice_files = [f"{prefix}_spice_{spice}.wav" for spice in SPICES] + [f"{prefix}_spice_generic.wav"]
    number_rows = []
    for number, spoken in sorted(NUMBER_WORDS.items()):
        row = {"Number": number, "Spoken_Text": spoken, "Suggested_Filename": f"{prefix}_num_{number}.wav",
               "Recording_Status": "PENDING", "Notes": "Reusable integer component for 1–999."}
        row["Recording_Status"] = keep_status(prior_numbers.get(str(number)),
            {k: row[k] for k in ("Spoken_Text", "Suggested_Filename")}, "Recording_Status")
        number_rows.append(row)

    for record in records:
        did, text = record["Dialogue_ID"], record["Original_Text"]
        variables = PLACEHOLDER.findall(text)
        unknown = set(variables) - {"{currency}", "{spiceName}", "{quantityLabel}", "{ruleReply}"} - {f"{{{n}}}" for n in NUMBER_FIELDS}
        if unknown:
            raise ValueError(f"{did}: new placeholder(s) {sorted(unknown)} need an audio plan")
        components = [c for c in PLACEHOLDER.split(text) if c]
        if not components or "".join(components) != text:
            raise ValueError(f"{did}: assembly lost source characters")
        sequence, row_manifest, row_static = [], [], []
        static_index = 0
        for order, part in enumerate(components, 1):
            p = part[1:-1] if PLACEHOLDER.fullmatch(part) else ""
            component = {"Dialogue_ID": did, "Segment_Order": order,
                         "Source_Text_Or_Placeholder": part, "Is_Required": "YES", "Notes": ""}
            if not p:
                if re.search(r"\w", part, re.UNICODE):
                    static_index += 1
                    filename = f"{prefix}_{did.rsplit('_', 1)[1]}_s{static_index:02d}.wav"
                    component.update(Component_Type="STATIC_CLIP", Playback_Reference=filename,
                                     Runtime_Resolution_Rule="Play clip; Source_Text retains exact spacing and punctuation.")
                    segment = {"Segment_ID": f"{did}_S{static_index:02d}", "Dialogue_ID": did,
                               "Segment_Order": order, "Text_To_Record": part.strip(),
                               "Suggested_Filename": filename,
                               "Voice_Direction": direction(prefix, record["Scenario"], record["Mood"], record["Patience"]),
                               "Scenario": record["Scenario"], "Recording_Status": "PENDING", "Notes": ""}
                    row_static.append(segment)
                else:
                    component.update(Component_Type="SEPARATOR", Playback_Reference="",
                                     Runtime_Resolution_Rule="No audio; retain exact whitespace/punctuation in text and timing metadata.")
            elif p in NUMBER_FIELDS:
                component.update(Component_Type="NUMBER_SEQUENCE", Playback_Reference=f"{prefix}_num_<component>.wav",
                                 Runtime_Resolution_Rule=f"Resolve {{{p}}} as integer 1–999: hundreds + tens/teens + units; concatenate existing number clips.")
                if p == "offeredPrice":
                    component["Notes"] = "Player-stated price has no 999 cap; larger values are unresolved."
                    unresolved.append((did, part, "Player offer can exceed 999; 36-clip library cannot pronounce it.",
                                       "Set a supported spoken range or approve larger-number clips/fallback before runtime audio."))
            elif p == "currency":
                component.update(Component_Type="VOCABULARY", Playback_Reference=f"{prefix}_currency_varahas.wav",
                                 Runtime_Resolution_Rule="Current provider always emits 'varahas'; use singular only after a provider change.")
            elif p == "spiceName":
                component.update(Component_Type="VOCABULARY", Playback_Reference="; ".join(spice_files),
                                 Runtime_Resolution_Rule="Choose the character-specific file matching runtime spiceName; 'spice' is fallback.")
            elif p == "quantityLabel":
                component.update(Component_Type="COMPOSITE_QUANTITY", Playback_Reference="; ".join(units + [f"{prefix}_quantity_this_lot.wav"]),
                                 Runtime_Resolution_Rule="Resolve leading integer via number library, then singular/plural unit. Runtime label also includes parenthetical metric text; its spoken treatment is undecided.",
                                 Notes="Main quantity/unit is planned; complete parenthetical speech is unresolved.")
                unresolved.append((did, part, "FormatTraditionalQuantity emits labels such as '2 Seers (~560g)' and '1 Veesai (~1.4 kg)'; the approved integer clips do not specify how to speak '~', g/kg, or decimals.",
                                   "Decide whether speech omits the parenthetical while subtitles retain it, or record a complete metric-suffix composition scheme."))
            else:  # ruleReply
                component.update(Component_Type="UNRESOLVED_EXPRESSION", Playback_Reference="",
                                 Runtime_Resolution_Rule="RuleBasedNPCBrain supplies a context-specific full sentence; no bounded clip recipe exists.")
                unresolved.append((did, part, "Rule reply can be an arbitrary generated sentence with a specific clarification question.",
                                   "Design bounded recorded responses or a safe text-only fallback while preserving specific questions."))
            if component.get("Component_Type") == "VOCABULARY":
                for filename in component["Playback_Reference"].split("; "):
                    if filename not in vocab_files:
                        raise ValueError(f"{did}: missing vocabulary {filename}")
            if component.get("Component_Type") == "COMPOSITE_QUANTITY":
                for filename in component["Playback_Reference"].split("; "):
                    if filename not in vocab_files:
                        raise ValueError(f"{did}: missing quantity vocabulary {filename}")
            row_manifest.append(component)
            sequence.append(f"{order}:{component['Component_Type']}:{component['Playback_Reference'] or 'metadata'}")
        assembly = " | ".join(sequence)
        valid_old = prior_overview.get(did, {})
        manifest_fields = ("Segment_Order", "Component_Type", "Source_Text_Or_Placeholder",
                           "Playback_Reference", "Runtime_Resolution_Rule", "Is_Required")
        valid = (valid_old.get("Original_Text") == text and
                 valid_old.get("Ordered_Assembly_Sequence") == assembly and
                 all(valid_old.get(field) == record[field] for field in ("Scenario", "Mood", "Patience")) and
                 [tuple(old.get(field) or "" for field in manifest_fields) for old in prior_manifest.get(did, [])] ==
                 [tuple(new.get(field) or "" for field in manifest_fields) for new in row_manifest])
        for segment in row_static:
            previous = prior_static.get(segment["Segment_ID"]) if valid else None
            segment["Recording_Status"] = keep_status(previous,
                {k: segment[k] for k in ("Dialogue_ID", "Segment_Order", "Text_To_Record",
                                          "Suggested_Filename", "Voice_Direction", "Scenario")},
                "Recording_Status")
        overview.append({"Dialogue_ID": did, "NPC": record["NPC"], "Scenario": record["Scenario"],
                         "Mood": record["Mood"], "Patience": record["Patience"], "Original_Text": text,
                         "Dynamic_Variables": "; ".join(dict.fromkeys(variables)) if variables else "None",
                         "Segment_Count": len(components), "Ordered_Assembly_Sequence": assembly,
                         "Status": keep_status(valid_old if valid else None, {"Original_Text": text,
                             "Ordered_Assembly_Sequence": assembly, "Scenario": record["Scenario"],
                             "Mood": record["Mood"], "Patience": record["Patience"]}, "Status"),
                         "Notes": "Intentionally inactive TimePressure scenario." if record["Scenario"] == "TimePressure" else
                             "Unresolved dynamic audio; see Unresolved_Audio_Cases.csv." if any(x[0] == did for x in unresolved) else ""})
        static.extend(row_static)
        manifest.extend(row_manifest)

    if len({r["Segment_ID"] for r in static}) != len(static):
        raise ValueError(f"{prefix}: duplicate static segment ID")
    filenames = [r["Suggested_Filename"] for r in static + number_rows + vocab]
    if len(set(filenames)) != len(filenames):
        raise ValueError(f"{prefix}: filename collision")
    grouped = defaultdict(list)
    for part in manifest:
        grouped[part["Dialogue_ID"]].append(part)
    for row in overview:
        pieces = grouped[row["Dialogue_ID"]]
        if "".join(p["Source_Text_Or_Placeholder"] for p in pieces) != row["Original_Text"]:
            raise ValueError(f"{row['Dialogue_ID']}: reconstruction mismatch")
        if [p["Segment_Order"] for p in pieces] != list(range(1, len(pieces) + 1)):
            raise ValueError(f"{row['Dialogue_ID']}: bad order")
    planned = len(static) + len(number_rows) + len(vocab)
    recorded = sum(r["Recording_Status"] in {"RECORDED", "QA_APPROVED", "APPROVED"}
                   for r in static + number_rows + vocab)
    summary = [
        {"Metric": "Total templates", "Value": len(overview), "Notes": "Source C# templates."},
        {"Metric": "Static templates", "Value": sum(r["Dynamic_Variables"] == "None" for r in overview), "Notes": "One full-sentence clip each."},
        {"Metric": "Dynamic templates", "Value": sum(r["Dynamic_Variables"] != "None" for r in overview), "Notes": "Split at every placeholder."},
        {"Metric": "Static clip count", "Value": len(static), "Notes": "Every recordable fixed-text occurrence."},
        {"Metric": "Number clip count", "Value": len(number_rows), "Notes": "1–999 component library; no zero in valid local trades."},
        {"Metric": "Dynamic vocabulary clip count", "Value": len(vocab), "Notes": "Character-specific reusable non-number clips."},
        {"Metric": "Total planned recordable clips", "Value": planned, "Notes": "Static + number + vocabulary; unresolved speech excluded."},
        {"Metric": "Unresolved expressions", "Value": len(unresolved), "Notes": "Dialogue-placeholder cases; see CSV."},
        {"Metric": "Recording progress", "Value": f"{recorded} / {planned}", "Notes": "Counts RECORDED/APPROVED/QA_APPROVED only."},
    ]
    return {"name": WORKBOOK_NAMES[records[0]["Prefix_Upper"]], "npc": records[0]["NPC"], "prefix": prefix,
            "sheets": {
                "Dialogue Overview": {"headers": OVERVIEW, "rows": overview},
                "Static Recording Segments": {"headers": STATIC, "rows": static},
                "Number Library": {"headers": NUMBERS, "rows": number_rows},
                "Dynamic Vocabulary": {"headers": VOCAB, "rows": vocab},
                "Assembly Manifest": {"headers": MANIFEST, "rows": manifest},
                "Recording Summary": {"headers": SUMMARY, "rows": summary},
            }, "unresolved": unresolved,
            "counts": {"templates": len(overview), "static_templates": summary[1]["Value"],
                       "dynamic_templates": summary[2]["Value"], "static_clips": len(static),
                       "number_clips": len(number_rows), "vocabulary_clips": len(vocab),
                       "planned_clips": planned, "unresolved_cases": len(unresolved),
                       "recorded_or_approved_clips": recorded}}


def run_builder(payload_path):
    node = os.environ.get("RECORDING_NODE", "node")
    modules = Path(os.environ.get("RECORDING_NODE_MODULES",
        Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"))
    if not (modules / "@oai" / "artifact-tool").exists():
        raise RuntimeError("@oai/artifact-tool unavailable; set RECORDING_NODE_MODULES to its node_modules directory")
    subprocess.run([node, str(OUT / "write_recording_workbooks.mjs"), str(payload_path), str(OUT), str(modules)], check=True)


def verify_exported(books, source):
    all_ids = []
    source_by_id = {row["Dialogue_ID"]: row["Original_Text"] for row in source}
    for book in books:
        wb = load_workbook(OUT / book["name"], read_only=True, data_only=True)
        try:
            if wb.sheetnames != list(book["sheets"]):
                raise ValueError(f"{book['name']}: missing or reordered tabs")
            actual = {}
            for name, expected in book["sheets"].items():
                cells = wb[name].iter_rows(values_only=True)
                header = next(cells)
                if list(header) != expected["headers"]:
                    raise ValueError(f"{book['name']}: {name} header mismatch")
                actual[name] = [dict(zip(header, row)) for row in cells]
                if len(actual[name]) != len(expected["rows"]):
                    raise ValueError(f"{book['name']}: {name} row-count mismatch")
            over = actual["Dialogue Overview"]
            manifest = actual["Assembly Manifest"]
            statics = actual["Static Recording Segments"]
            numbers = actual["Number Library"]
            vocab = actual["Dynamic Vocabulary"]
            by_dialogue = defaultdict(list)
            for part in manifest:
                by_dialogue[part["Dialogue_ID"]].append(part)
            for row in over:
                did = row["Dialogue_ID"]
                all_ids.append(did)
                parts = by_dialogue[did]
                if row["Original_Text"] != source_by_id[did] or \
                        "".join(part["Source_Text_Or_Placeholder"] for part in parts) != row["Original_Text"] or \
                        row["Segment_Count"] != len(parts):
                    raise ValueError(f"{did}: exported assembly differs from source")
            file_names = [row["Suggested_Filename"] for row in statics + numbers + vocab]
            if len(file_names) != len(set(file_names)):
                raise ValueError(f"{book['name']}: exported filename collision")
            static_files = {row["Suggested_Filename"] for row in statics}
            vocab_files = {row["Suggested_Filename"] for row in vocab}
            for part in manifest:
                if part["Component_Type"] == "STATIC_CLIP" and part["Playback_Reference"] not in static_files:
                    raise ValueError(f"{part['Dialogue_ID']}: missing static recording row")
                if part["Component_Type"] in ("VOCABULARY", "COMPOSITE_QUANTITY"):
                    if any(ref not in vocab_files for ref in part["Playback_Reference"].split("; ")):
                        raise ValueError(f"{part['Dialogue_ID']}: missing vocabulary recording row")
            if {row["Number"] for row in numbers} != set(NUMBER_WORDS):
                raise ValueError(f"{book['name']}: incomplete number library")
            for n in range(1, 1000):
                parts = []
                remainder = n
                if remainder >= 100:
                    parts.append((remainder // 100) * 100)
                    remainder %= 100
                if remainder >= 20:
                    parts.append((remainder // 10) * 10)
                    remainder %= 10
                if remainder:
                    parts.append(remainder)
                if any(part not in NUMBER_WORDS for part in parts):
                    raise ValueError(f"{book['name']}: cannot compose {n}")
        finally:
            wb.close()
    if len(all_ids) != 1047 or len(set(all_ids)) != 1047:
        raise ValueError("Exported workbooks have missing or duplicate dialogue IDs")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    source = inventory_from_source()
    by_prefix = defaultdict(list)
    for record in source:
        by_prefix[record["Prefix_Upper"]].append(record)
    books = [build_character(by_prefix[p]) for p in EXPECTED_COUNTS]
    unresolved = [case for book in books for case in book["unresolved"]]
    payload = {"books": [{k: v for k, v in book.items() if k != "unresolved"} for book in books]}
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".json", dir=OUT, delete=False) as temp:
        json.dump(payload, temp, ensure_ascii=False)
        payload_path = Path(temp.name)
    try:
        run_builder(payload_path)
    finally:
        payload_path.unlink(missing_ok=True)
    verify_exported(books, source)
    with (OUT / "Unresolved_Audio_Cases.csv").open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["Dialogue_ID", "Placeholder", "Reason", "Recommended_Later_Decision"])
        writer.writerows(unresolved)
    stats = {"source_templates": len(source), "source_csv_exact_match": True,
             "scenario_count": len(set(r["Scenario"] for r in source)),
             "placeholder_occurrences": dict(Counter(v for r in source for v in PLACEHOLDER.findall(r["Original_Text"]))),
             "per_npc": {book["npc"]: book["counts"] for book in books},
             "total_planned_clips": sum(book["counts"]["planned_clips"] for book in books),
             "total_unresolved_cases": len(unresolved),
             "unresolved_by_placeholder": dict(Counter(case[1] for case in unresolved)),
             "id_stability": "positional IDs from source order; insertions/reordering can shift later IDs"}
    (OUT / "Recording_Summary.json").write_text(json.dumps(stats, indent=2), encoding="utf-8")
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    main()
